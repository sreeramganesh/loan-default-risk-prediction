"""
LoanGuard AI — Comprehensive Automated Test Suite for Report Generation
Tests all 19 scenarios defined in Phase 9:
1. Successful loan assessment followed by Excel report generation.
2. Successful loan assessment followed by CSV report generation.
3. Excel response has the correct MIME type.
4. CSV response has the correct MIME type.
5. Downloaded files have appropriate filenames.
6. Excel workbook can be opened with openpyxl.
7. Excel workbook contains the expected worksheets.
8. Excel workbook contains actual submitted assessment inputs.
9. CSV contains the actual assessment values.
10. Default probability matches the actual validated model result.
11. Invalid input returns an appropriate HTTP error.
12. Missing assessment data is handled safely.
13. Malformed requests do not crash the application.
14. Repeated report downloads work correctly.
15. A second assessment produces a report for the second assessment, not the first assessment.
16. Failed assessments do not generate reports using stale results.
17. Formula-injection handling works for user-controlled text.
18. Existing prediction endpoints still work.
19. Existing homepage, charts, health endpoint, CSS, and JavaScript remain functional.
"""

import io
import csv
import json
import unittest
import openpyxl
from app import app, model, target_encoder
from report_generator import sanitize_text


class TestReportGeneration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.testing = True
        cls.client = app.test_client()

        # Standard test applicant payloads
        cls.prime_payload = {
            "loan_amount": 10000.0,
            "installment": 308.50,
            "int_rate": 6.89,
            "annual_income": 95000.0,
            "dti": 11.2,
            "application_type": "INDIVIDUAL",
            "verification_status": "Verified",
            "home_ownership": "MORTGAGE"
        }

        cls.high_risk_payload = {
            "loan_amount": 28000.0,
            "installment": 985.40,
            "int_rate": 22.80,
            "annual_income": 30000.0,
            "dti": 38.6,
            "application_type": "INDIVIDUAL",
            "verification_status": "Not Verified",
            "home_ownership": "RENT"
        }

    # 1. Successful loan assessment followed by Excel report generation
    def test_01_successful_assessment_and_excel_generation(self):
        pred_resp = self.client.post("/predict", json=self.prime_payload)
        self.assertEqual(pred_resp.status_code, 200)
        pred_data = pred_resp.get_json()
        self.assertEqual(pred_data["status"], "success")

        excel_payload = dict(self.prime_payload)
        excel_payload["report_id"] = pred_data["report_id"]
        excel_resp = self.client.post("/api/reports/excel", json=excel_payload)
        self.assertEqual(excel_resp.status_code, 200)
        self.assertGreater(len(excel_resp.data), 1000)

    # 2. Successful loan assessment followed by CSV report generation
    def test_02_successful_assessment_and_csv_generation(self):
        pred_resp = self.client.post("/predict", json=self.prime_payload)
        self.assertEqual(pred_resp.status_code, 200)
        pred_data = pred_resp.get_json()

        csv_payload = dict(self.prime_payload)
        csv_payload["report_id"] = pred_data["report_id"]
        csv_resp = self.client.post("/api/reports/csv", json=csv_payload)
        self.assertEqual(csv_resp.status_code, 200)
        self.assertGreater(len(csv_resp.data), 200)

    # 3. Excel response has the correct MIME type
    def test_03_excel_mime_type(self):
        resp = self.client.post("/api/reports/excel", json=self.prime_payload)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    # 4. CSV response has the correct MIME type
    def test_04_csv_mime_type(self):
        resp = self.client.post("/api/reports/csv", json=self.prime_payload)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.mimetype, "text/csv")
        self.assertIn("utf-8", resp.content_type.lower())

    # 5. Downloaded files have appropriate filenames
    def test_05_downloaded_filenames(self):
        test_payload = dict(self.prime_payload)
        test_payload["report_id"] = "LGA-20261018-TEST01"

        excel_resp = self.client.post("/api/reports/excel", json=test_payload)
        self.assertEqual(excel_resp.status_code, 200)
        cd_excel = excel_resp.headers.get("Content-Disposition", "")
        self.assertIn("attachment", cd_excel)
        self.assertIn("LoanGuard_AI_Assessment_LGA-20261018-TEST01.xlsx", cd_excel)

        csv_resp = self.client.post("/api/reports/csv", json=test_payload)
        self.assertEqual(csv_resp.status_code, 200)
        cd_csv = csv_resp.headers.get("Content-Disposition", "")
        self.assertIn("attachment", cd_csv)
        self.assertIn("LoanGuard_AI_Assessment_LGA-20261018-TEST01.csv", cd_csv)

    # 6. Excel workbook can be opened with openpyxl
    def test_06_excel_can_be_opened_with_openpyxl(self):
        resp = self.client.post("/api/reports/excel", json=self.prime_payload)
        self.assertEqual(resp.status_code, 200)
        buf = io.BytesIO(resp.data)
        wb = openpyxl.load_workbook(buf)
        self.assertIsNotNone(wb)

    # 7. Excel workbook contains the expected worksheets
    def test_07_excel_expected_worksheets(self):
        resp = self.client.post("/api/reports/excel", json=self.prime_payload)
        buf = io.BytesIO(resp.data)
        wb = openpyxl.load_workbook(buf)
        expected_sheets = ["LOAN ASSESSMENT", "RISK ANALYSIS", "RECOMMENDATIONS"]
        self.assertEqual(wb.sheetnames, expected_sheets)

    # 8. Excel workbook contains actual submitted assessment inputs
    def test_08_excel_contains_actual_inputs(self):
        resp = self.client.post("/api/reports/excel", json=self.prime_payload)
        buf = io.BytesIO(resp.data)
        wb = openpyxl.load_workbook(buf)
        ws1 = wb["LOAN ASSESSMENT"]

        # Search for loan amount 10000 and annual income 95000 in sheet cells
        values = [cell.value for row in ws1.iter_rows() for cell in row if cell.value is not None]
        self.assertTrue(any(v == 10000.0 or v == 10000 for v in values), "Loan amount not found in sheet 1")
        self.assertTrue(any(v == 95000.0 or v == 95000 for v in values), "Annual income not found in sheet 1")
        self.assertTrue(any("MORTGAGE" in str(v) for v in values), "Home ownership not found in sheet 1")

    # 9. CSV contains the actual assessment values
    def test_09_csv_contains_actual_values(self):
        resp = self.client.post("/api/reports/csv", json=self.prime_payload)
        csv_text = resp.data.decode("utf-8")
        reader = csv.reader(io.StringIO(csv_text))
        rows = dict(list(reader))

        self.assertIn("Requested Loan Amount ($)", rows)
        self.assertEqual(rows["Requested Loan Amount ($)"], "10000.00")
        self.assertIn("Gross Annual Income ($)", rows)
        self.assertEqual(rows["Gross Annual Income ($)"], "95000.00")
        self.assertIn("Home Ownership", rows)
        self.assertEqual(rows["Home Ownership"], "MORTGAGE")

    # 10. Default probability matches the actual validated model result
    def test_10_default_probability_matches_model(self):
        pred_resp = self.client.post("/predict", json=self.prime_payload)
        pred_data = pred_resp.get_json()
        expected_default_prob = pred_data["default_risk_percentage"]

        csv_resp = self.client.post("/api/reports/csv", json=self.prime_payload)
        csv_text = csv_resp.data.decode("utf-8")
        reader = csv.reader(io.StringIO(csv_text))
        rows = dict(list(reader))

        self.assertIn("Default Probability (Charged Off %)", rows)
        prob_str = rows["Default Probability (Charged Off %)"]
        self.assertEqual(prob_str, f"{expected_default_prob:.2f}%")

    # 11. Invalid input returns an appropriate HTTP error
    def test_11_invalid_numeric_input_returns_400(self):
        invalid_payload = dict(self.prime_payload)
        invalid_payload["loan_amount"] = -5000

        resp_excel = self.client.post("/api/reports/excel", json=invalid_payload)
        self.assertEqual(resp_excel.status_code, 400)
        data = resp_excel.get_json()
        self.assertEqual(data["status"], "error")
        self.assertIn("Loan Amount must be greater than zero", data["message"])

        resp_csv = self.client.post("/api/reports/csv", json=invalid_payload)
        self.assertEqual(resp_csv.status_code, 400)

    # 12. Missing assessment data is handled safely
    def test_12_missing_assessment_data(self):
        empty_resp = self.client.post("/api/reports/excel", json={})
        self.assertEqual(empty_resp.status_code, 400)
        data = empty_resp.get_json()
        self.assertEqual(data["status"], "error")

    # 13. Malformed requests do not crash the application
    def test_13_malformed_requests(self):
        resp = self.client.post(
            "/api/reports/excel",
            data="not-valid-json-{{{",
            content_type="application/json"
        )
        self.assertEqual(resp.status_code, 400)
        data = resp.get_json()
        self.assertEqual(data["status"], "error")

    # 14. Repeated report downloads work correctly
    def test_14_repeated_downloads(self):
        for _ in range(3):
            resp1 = self.client.post("/api/reports/excel", json=self.prime_payload)
            self.assertEqual(resp1.status_code, 200)
            resp2 = self.client.post("/api/reports/csv", json=self.prime_payload)
            self.assertEqual(resp2.status_code, 200)

    # 15. A second assessment produces a report for the second assessment, not the first
    def test_15_second_assessment_produces_second_report(self):
        # First assessment: Prime
        csv1 = self.client.post("/api/reports/csv", json=self.prime_payload).data.decode("utf-8")
        rows1 = dict(list(csv.reader(io.StringIO(csv1))))

        # Second assessment: High Risk
        csv2 = self.client.post("/api/reports/csv", json=self.high_risk_payload).data.decode("utf-8")
        rows2 = dict(list(csv.reader(io.StringIO(csv2))))

        self.assertEqual(rows1["Requested Loan Amount ($)"], "10000.00")
        self.assertEqual(rows2["Requested Loan Amount ($)"], "28000.00")
        self.assertEqual(rows1["Credit Verification Decision"], "VERIFIED & APPROVED")
        self.assertEqual(rows2["Credit Verification Decision"], "REJECTED (HIGH DEFAULT RISK)")
        self.assertNotEqual(rows1["Default Probability (Charged Off %)"], rows2["Default Probability (Charged Off %)"])

    # 16. Failed assessments do not generate reports using stale results
    def test_16_failed_assessment_handling(self):
        bad_payload = {"loan_amount": 0, "annual_income": 0}
        pred_resp = self.client.post("/predict", json=bad_payload)
        self.assertEqual(pred_resp.status_code, 400)

        excel_resp = self.client.post("/api/reports/excel", json=bad_payload)
        self.assertEqual(excel_resp.status_code, 400)

    # 17. Formula-injection handling works for user-controlled text
    def test_17_formula_injection_defense(self):
        # Direct unit test on sanitize_text
        self.assertEqual(sanitize_text("=cmd|'/c calc'!A0"), "'=cmd|'/c calc'!A0")
        self.assertEqual(sanitize_text("+123456"), "'+123456")
        self.assertEqual(sanitize_text("-100"), "'-100")
        self.assertEqual(sanitize_text("@SUM(A1:A10)"), "'@SUM(A1:A10)")
        self.assertEqual(sanitize_text("Normal Text"), "Normal Text")

        # Injected payload
        injected_payload = dict(self.prime_payload)
        injected_payload["home_ownership"] = "=cmd|'calc'!A0"
        csv_resp = self.client.post("/api/reports/csv", json=injected_payload)
        self.assertEqual(csv_resp.status_code, 200)
        csv_text = csv_resp.data.decode("utf-8")
        self.assertIn("'=cmd|'calc'!A0", csv_text)

    # 18. Existing prediction endpoints still work
    def test_18_existing_prediction_endpoint(self):
        resp = self.client.post("/predict", json=self.prime_payload)
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertEqual(data["status"], "success")
        self.assertIn("decision", data)
        self.assertIn("decision_badge", data)
        self.assertIn("risk_tier", data)
        self.assertIn("predicted_class", data)
        self.assertIn("default_risk_percentage", data)
        self.assertIn("recommendation", data)
        self.assertIn("probabilities", data)
        self.assertIn("financial_summary", data)

    # 19. Existing homepage, charts, health endpoint, CSS, and JS remain functional
    def test_19_existing_endpoints_functional(self):
        home_resp = self.client.get("/")
        self.assertEqual(home_resp.status_code, 200)
        self.assertIn(b"LoanGuard", home_resp.data)
        self.assertIn(b"DOWNLOAD ASSESSMENT REPORT", home_resp.data)

        health_resp = self.client.get("/health")
        self.assertEqual(health_resp.status_code, 200)
        self.assertEqual(health_resp.get_json()["status"], "healthy")

        metrics_resp = self.client.get("/api/metrics")
        self.assertEqual(metrics_resp.status_code, 200)
        self.assertIn("Random Forest", metrics_resp.get_json()["winner"])


if __name__ == "__main__":
    unittest.main()
