"""
LoanGuard AI — Report Generation Engine
Generates downloadable, audit-ready Excel (.xlsx) and CSV (.csv) reports
from server-validated loan underwriting assessments.
"""

import io
import csv
import uuid
from datetime import datetime, timezone
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ====================================================================
# Formula Injection Defense
# ====================================================================
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

def sanitize_text(val):
    """
    Sanitizes user-controlled text strings to prevent formula injection
    in CSV and Excel spreadsheets.
    """
    if val is None:
        return "Not available"
    val_str = str(val).strip()
    if not val_str:
        return "Not available"
    if val_str.startswith(FORMULA_PREFIXES):
        return "'" + val_str
    return val_str


# ====================================================================
# Core Assessment & Report Data Structure Builder
# ====================================================================
def validate_and_assess(raw_data, model, target_encoder, feature_encoders, feature_columns):
    """
    Validates applicant parameters on the server, executes model inference,
    calculates financial diagnostics, and returns a verified assessment payload.
    """
    if model is None or target_encoder is None:
        raise ValueError("Model artifacts are not loaded on server.")

    # 1. Parse and strictly validate numeric inputs
    try:
        loan_amount = float(raw_data.get("loan_amount", 0))
        installment = float(raw_data.get("installment", 0))
        int_rate = float(raw_data.get("int_rate", 0))
        annual_income = float(raw_data.get("annual_income", 0))
        dti = float(raw_data.get("dti", 0))
    except (ValueError, TypeError) as e:
        raise ValueError(f"Invalid numeric input provided: {e}")

    if loan_amount <= 0:
        raise ValueError("Loan Amount must be greater than zero.")
    if annual_income <= 0:
        raise ValueError("Annual Income must be greater than zero.")
    if installment <= 0:
        raise ValueError("Installment must be greater than zero.")
    if int_rate < 0 or int_rate > 100:
        raise ValueError("Interest rate must be between 0 and 100%.")
    if dti < 0 or dti > 150:
        raise ValueError("Debt-to-Income (DTI) ratio must be between 0 and 150%.")

    # Normalize percentage scales if entered as e.g. 12.5% instead of 0.125
    int_rate_normalized = int_rate / 100.0 if int_rate > 1.0 else int_rate
    dti_normalized = dti / 100.0 if dti > 1.0 else dti

    # 2. Parse and encode categorical features
    app_type_raw = str(raw_data.get("application_type", "INDIVIDUAL")).strip()
    ver_status_raw = str(raw_data.get("verification_status", "Not Verified")).strip()
    home_own_raw = str(raw_data.get("home_ownership", "RENT")).strip()

    def safe_encode(col_name, val):
        if col_name in feature_encoders:
            enc = feature_encoders[col_name]
            if val in enc.classes_:
                return int(enc.transform([val])[0])
        return 0

    app_type_enc = safe_encode("application_type", app_type_raw)
    ver_status_enc = safe_encode("verification_status", ver_status_raw)
    home_own_enc = safe_encode("home_ownership", home_own_raw)

    # 3. Model Prediction DataFrame
    input_data = {
        "loan_amount": [loan_amount],
        "installment": [installment],
        "int_rate": [int_rate_normalized],
        "annual_income": [annual_income],
        "dti": [dti_normalized],
        "application_type": [app_type_enc],
        "verification_status": [ver_status_enc],
        "home_ownership": [home_own_enc]
    }
    input_df = pd.DataFrame(input_data)[feature_columns]

    # Model inference
    pred_idx = model.predict(input_df)[0]
    prediction_label = target_encoder.inverse_transform([pred_idx])[0]

    probabilities = {}
    if hasattr(model, "predict_proba"):
        probs = model.predict_proba(input_df)[0]
        for cls_name, prob in zip(target_encoder.classes_, probs):
            probabilities[cls_name] = round(float(prob) * 100, 2)
    else:
        probabilities = {
            "Fully Paid": 85.0 if prediction_label == "Fully Paid" else 15.0,
            "Charged Off": 85.0 if prediction_label == "Charged Off" else 15.0,
            "Current": 0.0
        }

    default_prob = probabilities.get("Charged Off", 0.0)
    current_prob = probabilities.get("Current", 0.0)
    fully_paid_prob = probabilities.get("Fully Paid", 0.0)

    # Financial health metrics
    monthly_burden = (installment * 12 / annual_income * 100) if annual_income > 0 else 0.0
    loan_to_income = (loan_amount / annual_income * 100) if annual_income > 0 else 0.0

    # Decision Matrix (identical to app.py)
    if prediction_label == "Fully Paid" and default_prob < 30.0:
        decision = "VERIFIED & APPROVED"
        decision_badge = "success"
        risk_tier = "Low Risk"
        recommendation = (
            "Applicant displays a strong debt-to-income profile with minimal risk of charge-off. "
            "Recommended for automatic verification and loan origination."
        )
    elif prediction_label == "Charged Off" or default_prob >= 35.0:
        decision = "REJECTED (HIGH DEFAULT RISK)"
        decision_badge = "danger"
        risk_tier = "High Risk"
        recommendation = (
            "High probability of default detected. Key drivers include high debt-to-income burden "
            "and elevated interest tier. Verification rejected."
        )
    else:
        decision = "MANUAL UNDERWRITER REVIEW"
        decision_badge = "warning"
        risk_tier = "Moderate Risk"
        recommendation = (
            "Applicant falls into a borderline risk bracket. Secondary income verification, "
            "credit pull, and collateral assessment recommended before disbursement."
        )

    return {
        "status": "success",
        "loan_amount": loan_amount,
        "installment": installment,
        "int_rate_pct": int_rate_normalized * 100,
        "annual_income": annual_income,
        "monthly_income": annual_income / 12.0,
        "dti_pct": dti_normalized * 100,
        "application_type": app_type_raw,
        "verification_status": ver_status_raw,
        "home_ownership": home_own_raw,
        "decision": decision,
        "decision_badge": decision_badge,
        "risk_tier": risk_tier,
        "predicted_class": prediction_label,
        "default_risk_percentage": default_prob,
        "fully_paid_percentage": fully_paid_prob,
        "current_percentage": current_prob,
        "recommendation": recommendation,
        "monthly_burden_pct": monthly_burden,
        "loan_to_income_pct": loan_to_income
    }


def build_report_data(assessment, custom_report_id=None, custom_timestamp=None):
    """
    Constructs a comprehensive, immutable report dictionary partitioned
    into standard underwriter audit sections.
    """
    now_utc = datetime.now(timezone.utc)
    if custom_report_id and isinstance(custom_report_id, str) and len(custom_report_id.strip()) > 3:
        clean_id = "".join(c for c in custom_report_id.strip() if c.isalnum() or c in ("-", "_"))[:32]
        report_id = clean_id or f"LGA-{now_utc.strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"
    else:
        report_id = f"LGA-{now_utc.strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"

    if custom_timestamp and isinstance(custom_timestamp, str) and len(custom_timestamp.strip()) > 5:
        timestamp_str = sanitize_text(custom_timestamp.strip())[:40]
    else:
        timestamp_str = now_utc.strftime("%Y-%m-%d %H:%M:%S UTC")

    # Financial Diagnostic Evaluations
    dti = assessment["dti_pct"]
    if dti < 20.0:
        dti_status = "Excellent (< 20%)"
    elif dti <= 35.0:
        dti_status = "Moderate (20-35%)"
    else:
        dti_status = "Critical (> 35%)"

    burden = assessment["monthly_burden_pct"]
    if burden < 15.0:
        burden_status = "Manageable (< 15% of income)"
    elif burden <= 25.0:
        burden_status = "Elevated (15-25% of income)"
    else:
        burden_status = "High Burden (> 25% of income)"

    lti = assessment["loan_to_income_pct"]
    lti_multiple = f"{lti / 100.0:.2f}x"

    # Verification specific notes
    ver_status = assessment["verification_status"]
    if ver_status == "Verified":
        ver_note = "Income and employer documents have been independently verified."
    elif ver_status == "Source Verified":
        ver_note = "Income stream verified directly via automated bank feed or payroll source."
    else:
        ver_note = "Income stated by applicant. Proof of W-2 or tax returns recommended."

    return {
        "report_id": report_id,
        "timestamp": timestamp_str,
        "project_name": "LoanGuard AI - Autonomous Credit Verification Engine",
        "author": "S Sreeram Ganesh",
        "model_name": "Random Forest Classifier (Balanced)",
        "model_architecture": "Ensemble (n_estimators=200, criterion='gini', class_weight='balanced')",
        "benchmark_dataset": "Historical LendingClub Dataset (38,574 Samples)",
        
        # Loan Parameters
        "loan_amount": assessment["loan_amount"],
        "installment": assessment["installment"],
        "int_rate_pct": assessment["int_rate_pct"],
        "application_type": sanitize_text(assessment["application_type"]),
        "loan_purpose": "Personal / Debt Consolidation (Standard Credit Facility)",
        
        # Borrower Financials
        "annual_income": assessment["annual_income"],
        "monthly_income": assessment["monthly_income"],
        "dti_pct": assessment["dti_pct"],
        "monthly_burden_pct": assessment["monthly_burden_pct"],
        "loan_to_income_pct": assessment["loan_to_income_pct"],
        "dti_status": dti_status,
        "burden_status": burden_status,
        "lti_multiple": lti_multiple,

        # Verification & Collateral
        "home_ownership": sanitize_text(assessment["home_ownership"]),
        "verification_status": sanitize_text(assessment["verification_status"]),
        "verification_note": ver_note,

        # Machine Learning Assessment
        "decision": sanitize_text(assessment["decision"]),
        "decision_badge": assessment["decision_badge"],
        "risk_tier": sanitize_text(assessment["risk_tier"]),
        "predicted_class": sanitize_text(assessment["predicted_class"]),
        "default_risk_percentage": assessment["default_risk_percentage"],
        "fully_paid_percentage": assessment["fully_paid_percentage"],
        "current_percentage": assessment["current_percentage"],
        "recommendation": sanitize_text(assessment["recommendation"]),

        # Compliance & Limitations
        "disclaimer": (
            "DISCLAIMER: This automated assessment report is generated by LoanGuard AI for "
            "credit underwriting assistance and preliminary risk evaluation. It does not constitute "
            "a legally binding commitment to lend or a final credit denial. Final approval is subject "
            "to institutional underwriting policy, secondary documentation review, and compliance "
            "with the Equal Credit Opportunity Act (ECOA) and applicable fair lending regulations."
        ),
        "limitations": (
            "Model predictions are computed using 8 pre-origination attributes trained on historical "
            "lending outcomes. The model does not monitor post-disbursement macroeconomic events, "
            "applicant employment loss, or sudden external economic disruptions."
        )
    }


# ====================================================================
# Excel Report Generation (openpyxl)
# ====================================================================
def generate_excel_report(report_data):
    """
    Generates a beautifully styled 3-sheet Excel workbook in-memory.
    Returns: io.BytesIO buffer containing the XLSX file.
    """
    wb = openpyxl.Workbook()
    
    # Define Palette & Styles
    NAVY_HEADER_FILL = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
    SECTION_FILL = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
    SUB_HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Status Fills
    badge = report_data.get("decision_badge", "warning")
    if badge == "success":
        STATUS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        STATUS_FONT = Font(name="Calibri", size=11, bold=True, color="166534")
    elif badge == "danger":
        STATUS_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        STATUS_FONT = Font(name="Calibri", size=11, bold=True, color="991B1B")
    else:
        STATUS_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        STATUS_FONT = Font(name="Calibri", size=11, bold=True, color="92400E")

    FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="38BDF8")
    FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="CBD5E1")
    FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_TH = Font(name="Calibri", size=10, bold=True, color="E2E8F0")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
    FONT_REGULAR = Font(name="Calibri", size=10, color="1E293B")
    FONT_MUTED = Font(name="Calibri", size=9, italic=True, color="64748B")

    THIN_SIDE = Side(border_style="thin", color="E2E8F0")
    BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
    BORDER_BOTTOM = Border(bottom=Side(border_style="medium", color="0284C7"))

    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ----------------------------------------------------------------
    # SHEET 1: LOAN ASSESSMENT
    # ----------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "LOAN ASSESSMENT"
    ws1.views.sheetView[0].showGridLines = True

    # Title Block
    ws1.merge_cells("A1:D1")
    ws1["A1"] = "LOANGUARD AI — CREDIT UNDERWRITING ASSESSMENT REPORT"
    ws1["A1"].font = FONT_TITLE
    ws1["A1"].fill = NAVY_HEADER_FILL
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:D2")
    ws1["A2"] = f"Report ID: {report_data['report_id']}   |   Assessment Date: {report_data['timestamp']}   |   Model: {report_data['model_name']}"
    ws1["A2"].font = FONT_SUBTITLE
    ws1["A2"].fill = SECTION_FILL
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    row_idx = 4

    def add_section_header(ws, row, title):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        c = ws.cell(row=row, column=1, value=title)
        c.font = FONT_SECTION
        c.fill = SUB_HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = 24
        return row + 1

    def add_table_header(ws, row, col1_title, col2_title, col3_title=None, col4_title=None):
        ws.cell(row=row, column=1, value=col1_title).font = FONT_TH
        ws.cell(row=row, column=1).fill = SECTION_FILL
        ws.cell(row=row, column=1).alignment = ALIGN_LEFT
        ws.cell(row=row, column=2, value=col2_title).font = FONT_TH
        ws.cell(row=row, column=2).fill = SECTION_FILL
        ws.cell(row=row, column=2).alignment = ALIGN_RIGHT

        if col3_title:
            ws.cell(row=row, column=3, value=col3_title).font = FONT_TH
            ws.cell(row=row, column=3).fill = SECTION_FILL
            ws.cell(row=row, column=3).alignment = ALIGN_LEFT
        if col4_title:
            ws.cell(row=row, column=4, value=col4_title).font = FONT_TH
            ws.cell(row=row, column=4).fill = SECTION_FILL
            ws.cell(row=row, column=4).alignment = ALIGN_RIGHT
        
        ws.row_dimensions[row].height = 20
        return row + 1

    # Section A: Primary Decision Summary
    row_idx = add_section_header(ws1, row_idx, "01. EXECUTIVE UNDERWRITING DECISION")
    
    # Decision Card
    ws1.cell(row=row_idx, column=1, value="Verification Decision:").font = FONT_BOLD
    ws1.cell(row=row_idx, column=1).alignment = ALIGN_LEFT
    c_dec = ws1.cell(row=row_idx, column=2, value=report_data["decision"])
    c_dec.font = STATUS_FONT
    c_dec.fill = STATUS_FILL
    c_dec.alignment = ALIGN_CENTER

    ws1.cell(row=row_idx, column=3, value="Risk Tier:").font = FONT_BOLD
    ws1.cell(row=row_idx, column=3).alignment = ALIGN_LEFT
    c_tier = ws1.cell(row=row_idx, column=4, value=f"{report_data['risk_tier']} ({report_data['default_risk_percentage']:.1f}% Default Risk)")
    c_tier.font = STATUS_FONT
    c_tier.fill = STATUS_FILL
    c_tier.alignment = ALIGN_CENTER
    ws1.row_dimensions[row_idx].height = 22
    row_idx += 1

    ws1.cell(row=row_idx, column=1, value="Predicted Outcome:").font = FONT_BOLD
    ws1.cell(row=row_idx, column=2, value=report_data["predicted_class"]).font = FONT_REGULAR
    ws1.cell(row=row_idx, column=3, value="Default Probability:").font = FONT_BOLD
    c_def = ws1.cell(row=row_idx, column=4, value=report_data["default_risk_percentage"] / 100.0)
    c_def.font = FONT_BOLD
    c_def.number_format = "0.0%"
    c_def.alignment = ALIGN_RIGHT
    row_idx += 2

    # Section B: Loan Terms & Facilities
    row_idx = add_section_header(ws1, row_idx, "02. REQUESTED LOAN TERMS")
    loan_items = [
        ("Requested Loan Amount", report_data["loan_amount"], "$#,##0.00"),
        ("Monthly Installment", report_data["installment"], "$#,##0.00"),
        ("Interest Rate (APR)", report_data["int_rate_pct"] / 100.0, "0.00%"),
        ("Stated Loan Purpose", report_data["loan_purpose"], None),
        ("Application Type", report_data["application_type"], None)
    ]
    for label, val, num_fmt in loan_items:
        ws1.cell(row=row_idx, column=1, value=label).font = FONT_BOLD
        c = ws1.cell(row=row_idx, column=2, value=val)
        c.font = FONT_REGULAR
        if num_fmt:
            c.number_format = num_fmt
            c.alignment = ALIGN_RIGHT
        else:
            c.alignment = ALIGN_LEFT
        ws1.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        ws1.row_dimensions[row_idx].height = 19
        row_idx += 1
    row_idx += 1

    # Section C: Borrower Financial Inputs & Ratios
    row_idx = add_section_header(ws1, row_idx, "03. BORROWER FINANCIAL PROFILE")
    fin_items = [
        ("Gross Annual Income", report_data["annual_income"], "$#,##0.00"),
        ("Gross Monthly Income", report_data["monthly_income"], "$#,##0.00"),
        ("Debt-to-Income Ratio (DTI)", report_data["dti_pct"] / 100.0, "0.0%"),
        ("Estimated Monthly Debt Burden", report_data["monthly_burden_pct"] / 100.0, "0.0%"),
        ("Loan-to-Income Multiple", report_data["lti_multiple"], None),
        ("DTI Evaluation Tier", report_data["dti_status"], None),
        ("Debt Burden Evaluation", report_data["burden_status"], None)
    ]
    for label, val, num_fmt in fin_items:
        ws1.cell(row=row_idx, column=1, value=label).font = FONT_BOLD
        c = ws1.cell(row=row_idx, column=2, value=val)
        c.font = FONT_REGULAR
        if num_fmt:
            c.number_format = num_fmt
            c.alignment = ALIGN_RIGHT
        else:
            c.alignment = ALIGN_LEFT
        ws1.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        ws1.row_dimensions[row_idx].height = 19
        row_idx += 1
    row_idx += 1

    # Section D: Verification & Collateral
    row_idx = add_section_header(ws1, row_idx, "04. VERIFICATION & COLLATERAL")
    ver_items = [
        ("Home Ownership Status", report_data["home_ownership"]),
        ("Income Verification Status", report_data["verification_status"]),
        ("Underwriter Verification Note", report_data["verification_note"])
    ]
    for label, val in ver_items:
        ws1.cell(row=row_idx, column=1, value=label).font = FONT_BOLD
        c = ws1.cell(row=row_idx, column=2, value=val)
        c.font = FONT_REGULAR
        c.alignment = ALIGN_LEFT
        ws1.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        ws1.row_dimensions[row_idx].height = 19
        row_idx += 1

    # Set column widths for Sheet 1
    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 24
    ws1.column_dimensions["C"].width = 22
    ws1.column_dimensions["D"].width = 26

    # ----------------------------------------------------------------
    # SHEET 2: RISK ANALYSIS
    # ----------------------------------------------------------------
    ws2 = wb.create_sheet(title="RISK ANALYSIS")
    ws2.views.sheetView[0].showGridLines = True

    # Title Block
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "LOANGUARD AI — MULTICLASS RISK ANALYSIS & DIAGNOSTICS"
    ws2["A1"].font = FONT_TITLE
    ws2["A1"].fill = NAVY_HEADER_FILL
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:D2")
    ws2["A2"] = f"Report ID: {report_data['report_id']}   |   Decision: {report_data['decision']} ({report_data['risk_tier']})"
    ws2["A2"].font = FONT_SUBTITLE
    ws2["A2"].fill = SECTION_FILL
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 20

    row_idx = 4
    row_idx = add_section_header(ws2, row_idx, "01. MULTICLASS CONFIDENCE DISTRIBUTION")
    row_idx = add_table_header(ws2, row_idx, "Target Loan Class", "Model Probability", "Business Meaning", "Risk Classification")

    prob_rows = [
        ("Fully Paid", report_data["fully_paid_percentage"] / 100.0, "Good Standing / Full Amortization", "Safe / Prime"),
        ("Charged Off", report_data["default_risk_percentage"] / 100.0, "Credit Default / Severe Lender Loss", "High Default Risk"),
        ("Current", report_data["current_percentage"] / 100.0, "Active Payments / In Grace Period", "Neutral / Ongoing")
    ]
    for cls_name, prob, meaning, risk_class in prob_rows:
        ws2.cell(row=row_idx, column=1, value=cls_name).font = FONT_BOLD
        c_p = ws2.cell(row=row_idx, column=2, value=prob)
        c_p.font = FONT_BOLD
        c_p.number_format = "0.0%"
        c_p.alignment = ALIGN_RIGHT

        ws2.cell(row=row_idx, column=3, value=meaning).font = FONT_REGULAR
        ws2.cell(row=row_idx, column=4, value=risk_class).font = FONT_REGULAR
        ws2.row_dimensions[row_idx].height = 20
        row_idx += 1
    row_idx += 1

    # Section 2: Established Policy Thresholds
    row_idx = add_section_header(ws2, row_idx, "02. UNDERWRITING RISK THRESHOLD BENCHMARKS")
    row_idx = add_table_header(ws2, row_idx, "Risk Bracket", "Default Prob Range", "Underwriting Action", "Current Applicant Match")

    thresh_rows = [
        ("Prime / Low Risk", "< 30.0%", "Auto-Verify & Originate", "MATCH" if report_data["risk_tier"] == "Low Risk" else "No"),
        ("Near-Prime / Review", "30.0% – 34.9%", "Manual Underwriter Review", "MATCH" if report_data["risk_tier"] == "Moderate Risk" else "No"),
        ("Subprime / High Risk", ">= 35.0%", "Automatic Rejection", "MATCH" if report_data["risk_tier"] == "High Risk" else "No")
    ]
    for bracket, p_range, action, match in thresh_rows:
        ws2.cell(row=row_idx, column=1, value=bracket).font = FONT_BOLD
        ws2.cell(row=row_idx, column=2, value=p_range).font = FONT_REGULAR
        ws2.cell(row=row_idx, column=2).alignment = ALIGN_CENTER
        ws2.cell(row=row_idx, column=3, value=action).font = FONT_REGULAR

        c_m = ws2.cell(row=row_idx, column=4, value=match)
        if match == "MATCH":
            c_m.font = STATUS_FONT
            c_m.fill = STATUS_FILL
            c_m.alignment = ALIGN_CENTER
        else:
            c_m.font = FONT_MUTED
            c_m.alignment = ALIGN_CENTER
        ws2.row_dimensions[row_idx].height = 20
        row_idx += 1
    row_idx += 1

    # Section 3: Diagnostic Financial Health Indicators
    row_idx = add_section_header(ws2, row_idx, "03. FINANCIAL RATIO DIAGNOSTICS")
    diag_rows = [
        ("Monthly Debt Burden", f"{report_data['monthly_burden_pct']:.1f}%", report_data["burden_status"], "Benchmarked < 15%"),
        ("Debt-to-Income (DTI)", f"{report_data['dti_pct']:.1f}%", report_data["dti_status"], "Industry Standard < 36%"),
        ("Loan-to-Income Leverage", report_data["lti_multiple"], "Leverage multiple against gross income", "Target < 0.35x")
    ]
    for label, val_str, status, benchmark in diag_rows:
        ws2.cell(row=row_idx, column=1, value=label).font = FONT_BOLD
        ws2.cell(row=row_idx, column=2, value=val_str).font = FONT_BOLD
        ws2.cell(row=row_idx, column=2).alignment = ALIGN_RIGHT
        ws2.cell(row=row_idx, column=3, value=status).font = FONT_REGULAR
        ws2.cell(row=row_idx, column=4, value=benchmark).font = FONT_MUTED
        ws2.row_dimensions[row_idx].height = 20
        row_idx += 1

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 36
    ws2.column_dimensions["D"].width = 28

    # ----------------------------------------------------------------
    # SHEET 3: RECOMMENDATIONS
    # ----------------------------------------------------------------
    ws3 = wb.create_sheet(title="RECOMMENDATIONS")
    ws3.views.sheetView[0].showGridLines = True

    # Title Block
    ws3.merge_cells("A1:D1")
    ws3["A1"] = "LOANGUARD AI — UNDERWRITER GUIDANCE & COMPLIANCE"
    ws3["A1"].font = FONT_TITLE
    ws3["A1"].fill = NAVY_HEADER_FILL
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 36

    ws3.merge_cells("A2:D2")
    ws3["A2"] = f"Report ID: {report_data['report_id']}   |   Status: {report_data['decision']}"
    ws3["A2"].font = FONT_SUBTITLE
    ws3["A2"].fill = SECTION_FILL
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[2].height = 20

    row_idx = 4

    # Guidance Section
    row_idx = add_section_header(ws3, row_idx, "01. ACTIONABLE UNDERWRITER GUIDANCE")
    ws3.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx + 1, end_column=4)
    c_rec = ws3.cell(row=row_idx, column=1, value=report_data["recommendation"])
    c_rec.font = FONT_BOLD
    c_rec.alignment = ALIGN_WRAP
    c_rec.fill = ZEBRA_FILL
    ws3.row_dimensions[row_idx].height = 25
    ws3.row_dimensions[row_idx + 1].height = 25
    row_idx += 3

    # Verification Checklist
    row_idx = add_section_header(ws3, row_idx, "02. UNDERWRITING VERIFICATION CHECKLIST")
    checklist = [
        ("Income Verification", f"Current Status: {report_data['verification_status']}", report_data["verification_note"]),
        ("Collateral / Housing", f"Housing Status: {report_data['home_ownership']}", "Confirm mortgage statement or active lease agreement."),
        ("Debt Obligation Review", f"DTI: {report_data['dti_pct']:.1f}% ({report_data['dti_status']})", "Check credit bureau liabilities against reported debt burden."),
        ("Identity Verification", "Application Type: " + report_data["application_type"], "Perform CIP/KYC identity and OFAC sanctions screen.")
    ]
    for item, stat, note in checklist:
        ws3.cell(row=row_idx, column=1, value=item).font = FONT_BOLD
        ws3.cell(row=row_idx, column=2, value=stat).font = FONT_REGULAR
        ws3.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
        c_n = ws3.cell(row=row_idx, column=3, value=note)
        c_n.font = FONT_REGULAR
        c_n.alignment = ALIGN_LEFT
        ws3.row_dimensions[row_idx].height = 20
        row_idx += 1
    row_idx += 1

    # Limitations & Assumptions
    row_idx = add_section_header(ws3, row_idx, "03. MODEL ASSUMPTIONS & LIMITATIONS")
    ws3.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx + 1, end_column=4)
    c_lim = ws3.cell(row=row_idx, column=1, value=report_data["limitations"])
    c_lim.font = FONT_MUTED
    c_lim.alignment = ALIGN_WRAP
    c_lim.fill = ZEBRA_FILL
    ws3.row_dimensions[row_idx].height = 24
    ws3.row_dimensions[row_idx + 1].height = 24
    row_idx += 3

    # Legal Disclaimer
    row_idx = add_section_header(ws3, row_idx, "04. REGULATORY & FAIR LENDING DISCLAIMER")
    ws3.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx + 2, end_column=4)
    c_disc = ws3.cell(row=row_idx, column=1, value=report_data["disclaimer"])
    c_disc.font = FONT_MUTED
    c_disc.alignment = ALIGN_WRAP
    c_disc.fill = ZEBRA_FILL
    ws3.row_dimensions[row_idx].height = 22
    ws3.row_dimensions[row_idx + 1].height = 22
    ws3.row_dimensions[row_idx + 2].height = 22

    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 30
    ws3.column_dimensions["D"].width = 30

    # Save to BytesIO
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer


# ====================================================================
# CSV Report Generation (csv module)
# ====================================================================
def generate_csv_report(report_data):
    """
    Generates a clean, UTF-8 two-column Field,Value CSV report in-memory.
    Returns: io.BytesIO buffer containing the CSV file.
    """
    csv_rows = [
        # Report Metadata
        ("Report ID", report_data["report_id"]),
        ("Assessment Timestamp", report_data["timestamp"]),
        ("Project Name", report_data["project_name"]),
        ("Author", report_data["author"]),
        ("Model Name", report_data["model_name"]),
        ("Model Architecture", report_data["model_architecture"]),
        ("Training Dataset", report_data["benchmark_dataset"]),
        
        # Loan Parameters
        ("Requested Loan Amount ($)", f"{report_data['loan_amount']:.2f}"),
        ("Monthly Installment ($/mo)", f"{report_data['installment']:.2f}"),
        ("Interest Rate (%)", f"{report_data['int_rate_pct']:.2f}%"),
        ("Stated Loan Purpose", report_data["loan_purpose"]),
        ("Application Type", report_data["application_type"]),

        # Borrower Financial Profile
        ("Gross Annual Income ($)", f"{report_data['annual_income']:.2f}"),
        ("Gross Monthly Income ($)", f"{report_data['monthly_income']:.2f}"),
        ("Debt-to-Income Ratio (DTI %)", f"{report_data['dti_pct']:.2f}%"),
        ("Estimated Monthly Debt Burden (%)", f"{report_data['monthly_burden_pct']:.2f}%"),
        ("Loan-to-Income Ratio (%)", f"{report_data['loan_to_income_pct']:.2f}%"),
        ("Loan-to-Income Multiple", report_data["lti_multiple"]),
        ("DTI Evaluation Tier", report_data["dti_status"]),
        ("Debt Burden Evaluation", report_data["burden_status"]),

        # Verification & Collateral
        ("Home Ownership", report_data["home_ownership"]),
        ("Income Verification Status", report_data["verification_status"]),
        ("Verification Note", report_data["verification_note"]),

        # Machine Learning Assessment
        ("Credit Verification Decision", report_data["decision"]),
        ("Risk Tier", report_data["risk_tier"]),
        ("Predicted Class", report_data["predicted_class"]),
        ("Default Probability (Charged Off %)", f"{report_data['default_risk_percentage']:.2f}%"),
        ("Good Standing Probability (Fully Paid %)", f"{report_data['fully_paid_percentage']:.2f}%"),
        ("Active Grace Probability (Current %)", f"{report_data['current_percentage']:.2f}%"),
        
        # Summary & Guidance
        ("Actionable Underwriter Guidance", report_data["recommendation"]),
        ("Model Assumptions & Limitations", report_data["limitations"]),
        ("Regulatory & Fair Lending Disclaimer", report_data["disclaimer"])
    ]

    string_buffer = io.StringIO()
    writer = csv.writer(string_buffer, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(["Field", "Value"])

    for field, val in csv_rows:
        safe_val = sanitize_text(val)
        writer.writerow([field, safe_val])

    # Convert to UTF-8 BytesIO
    csv_bytes = string_buffer.getvalue().encode("utf-8")
    bytes_buffer = io.BytesIO(csv_bytes)
    bytes_buffer.seek(0)
    return bytes_buffer
